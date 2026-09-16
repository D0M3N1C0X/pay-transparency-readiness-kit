"""The Excel model reproduces the pandas pipeline, formula by formula.

Evaluating the full workbook takes minutes in pure Python, so this test builds a sample
workbook and calculates it with the `formulas` package. CI also recalculates the full
workbook with LibreOffice (see .github/workflows/ci.yml).
"""
import pandas as pd
import pytest

formulas = pytest.importorskip("formulas")

import build_workbook  # noqa: E402
import check_workbook  # noqa: E402
from config import PAYROLL_EXTRACT  # noqa: E402
from workers import derive  # noqa: E402


@pytest.fixture(scope="module")
def sample_model(tmp_path_factory):
    extract = pd.read_csv(PAYROLL_EXTRACT)
    sample = pd.concat([g.sample(n=25, random_state=11) for _, g in extract.groupby("country")])
    path = tmp_path_factory.mktemp("wb") / "sample.xlsx"
    model = build_workbook.build(derive(sample.sort_values("worker_id")), path)
    return model, path


def test_every_reconciliation_row_matches_and_nothing_errors(sample_model, capsys):
    model, path = sample_model
    values = check_workbook.values_from_formulas(path)
    assert check_workbook.check(values, model.wb.sheetnames, len(model.checks)) == 0, capsys.readouterr().out


def test_a_wrong_python_value_is_caught(sample_model, tmp_path):
    from openpyxl import load_workbook
    model, path = sample_model
    wb = load_workbook(path)
    ws = wb["Reconciliation"]
    row = next(r for r in range(7, 7 + len(model.checks)) if isinstance(ws[f"D{r}"].value, float))
    ws[f"D{row}"].value += 1e-6
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    values = check_workbook.values_from_formulas(broken)
    assert check_workbook.check(values, wb.sheetnames, len(model.checks)) == 1
