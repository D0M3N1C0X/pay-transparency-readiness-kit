"""
The catalogue of reporting items: one card per Article 9(1) item and one for the Article 10 screen.

docs/method.md explains how each figure is computed. This page adds what a reviewer asks next: the result for
each employer, the cell and formula that produce it in the Excel model, the choice the kit made where the
Directive leaves room, and what a statutory report would need on top. Results and formulas are read from the
code and the workbook, so the page cannot fall behind; a test checks it is current.

    python src/build_catalogue.py
"""
from openpyxl import load_workbook

import article9
import article10
import config as C
from build_report import pct
from config import BANDS
from workers import load_workers

OUT = C.ROOT / "docs" / "catalogue.md"
MODEL = C.DELIVERABLES / "pay_transparency_model.xlsx"

CARDS = [
    {
        "title": "(a) Gender pay gap",
        "article": "Article 9(1)(a); definition in Article 3(1)(c)",
        "keys": ["a_mean_gap_hourly", "a_mean_gap_annual"],
        "choice": "Reported on hourly and on annual pay, because Article 3(1)(b) defines pay level as both. Pay is "
                  "base plus every component, annualised for the workers present at the snapshot.",
        "real": "Pay actually paid in the previous calendar year, with leavers and joiners treated as the national "
                "transposition says.",
    },
    {
        "title": "(b) Gap in complementary or variable components",
        "article": "Article 9(1)(b)",
        "keys": ["b_mean_gap_components"],
        "choice": "Computed among workers who receive any component. Over all workers it would mix the size of the "
                  "components with who is eligible for them, which item (e) already reports.",
        "real": "The national rule may fix the population; the choice must be stated in the report either way.",
    },
    {
        "title": "(c) Median gender pay gap",
        "article": "Article 9(1)(c)",
        "keys": ["c_median_gap_hourly", "c_median_gap_annual"],
        "choice": "The same formula as (a) on medians, hourly and annual.",
        "real": "As for (a).",
    },
    {
        "title": "(d) Median gap in complementary or variable components",
        "article": "Article 9(1)(d)",
        "keys": ["d_median_gap_components"],
        "choice": "Among recipients, as for (b).",
        "real": "As for (b).",
    },
    {
        "title": "(e) Share receiving components",
        "article": "Article 9(1)(e)",
        "keys": ["e_share_women_receiving", "e_share_men_receiving"],
        "choice": "A worker receives components if the sum of bonus, commission, shift and car allowance is above "
                  "zero.",
        "real": "The list of components comes from the payroll's own pay codes, mapped once with the client.",
    },
    {
        "title": "(f) Quartile pay bands",
        "article": "Article 9(1)(f)",
        "keys": [f"f_q{q}_share_women" for q in (1, 2, 3, 4)],
        "choice": "Workers ranked by hourly total pay within each employer and cut into four groups of equal size. "
                  "Equal pay at a boundary is split by the order of the payroll extract, so the bands stay equal in "
                  "size. The men's shares are the complement and are in the workbook.",
        "real": "National law may define the ranking pay (hourly or annual) and how ties are split.",
    },
]

# rows of the "Article 9" sheet hold the items; the workings they point to are in rows 30-56
ITEM_ROWS = {
    "a_mean_gap_hourly": 9, "a_mean_gap_annual": 10, "b_mean_gap_components": 11, "c_median_gap_hourly": 12,
    "c_median_gap_annual": 13, "d_median_gap_components": 14, "e_share_women_receiving": 15,
    "e_share_men_receiving": 16, **{f"f_q{q}_share_women": 16 + q for q in (1, 2, 3, 4)},
}


def _formula(cell) -> str:
    v = cell.value
    return getattr(v, "text", v) if v is not None else ""


def _results(ind, keys: list[str]) -> list[str]:
    cols = [*C.ENTITIES, "Group"]
    rows = ["| Item | " + " | ".join(cols) + " |", "|---|" + "---:|" * len(cols)]
    labels = {k: label for k, label, _ in article9.ITEMS}
    for k in keys:
        rows.append(f"| {labels[k]} | " + " | ".join(pct(ind.loc[k, c]) for c in cols) + " |")
    return rows


def card(c: dict, ind, ws) -> list[str]:
    lines = [f"## {c['title']}", "", f"**Directive.** {c['article']}. How it is computed: "
             f"[method.md](method.md#the-article-91-items).", "", "**Result at 30 June 2026.**", "",
             *_results(ind, c["keys"]), "",
             "**In the Excel model** (sheet Article 9, column C for the Italian employer; the other employers are the "
             "next columns):", "", "```"]
    for k in c["keys"]:
        r = ITEM_ROWS[k]
        lines.append(f"C{r}  {ws[f'B{r}'].value}")
        lines.append(f"     {_formula(ws[f'C{r}'])}")
    lines += ["```", "", f"**The kit's choice.** {c['choice']}", "",
              f"**In a statutory report.** {c['real']}", ""]
    return lines


def screen_card(screen) -> list[str]:
    flagged = screen[screen["status"] == "Justify or remedy"]
    by_country = flagged.groupby("country").size()
    rows = ["| Employer | Categories | Flagged | Upper-bound cost to close, a year |", "|---|---:|---:|---:|"]
    for code, name in C.ENTITIES.items():
        n = int((screen["country"] == code).sum())
        cost = flagged.loc[flagged["country"] == code, "remedy_cost"].sum()
        rows.append(f"| {name} | {n} | {int(by_country.get(code, 0))} | €{cost:,.0f} |")
    total = flagged["remedy_cost"].sum()
    rows.append(f"| **Group** | **{len(screen)}** | **{len(flagged)}** | **€{total:,.0f}** |")
    return ["## (g) Categories of workers, and the Article 10 screen", "",
            "**Directive.** Article 9(1)(g) for the gap by category; Article 10(1) for the joint pay assessment, owed "
            "when a difference of at least 5% in a category is neither justified on objective, gender-neutral "
            "criteria nor remedied within six months. How it is computed: [method.md](method.md#the-article-10-screen).",
            "", f"**Result at 30 June 2026.** {len(flagged)} of {len(screen)} categories show a difference of "
            f"{pct(C.JPA_THRESHOLD, 0)} or more, on hourly or annual pay, and need a justification or a remedy:", "",
            *rows, "",
            "**In the Excel model.** Sheet Categories: one row per employer and category, with the gaps, the flag and "
            "two columns the client fills in, the justification and whether it was remedied.", "",
            "**The kit's choice.** Categories come from a gender-neutral job evaluation on the four Article 4(4) "
            "criteria, scored per role and banded into points; a flag is raised by the hourly or the annual gap; "
            "the cost is a ceiling, the mean hourly gap times the hours of the lower-paid sex.", "",
            "**In a statutory report.** Categories resting on the employer's own gender-neutral job evaluation, "
            "justifications documented on objective criteria before the report date, and any joint pay assessment "
            "carried out in cooperation with workers' representatives, as Article 10 requires.", ""]


def text() -> str:
    w = load_workers()
    codes = list(C.ENTITIES)
    ind = article9.by_entity(w, codes)
    screen = article10.screen(w, codes, [b for _, b in BANDS])
    ws = load_workbook(MODEL)["Article 9"]
    lines = ["# Catalogue of reporting items", "",
             "One card per Article 9(1) item and one for the Article 10 screen: the result for each employer, where "
             "the Excel model computes it, the choice the kit made and what a statutory report would add. Generated "
             "by `src/build_catalogue.py` from the code and the workbook, and checked by a test.", ""]
    for c in CARDS:
        lines += card(c, ind, ws)
    lines += screen_card(screen)
    return "\n".join(lines).rstrip("\n") + "\n"


def write() -> None:
    OUT.write_text(text(), encoding="utf-8")


if __name__ == "__main__":
    write()
    print(f"catalogue -> {OUT.relative_to(C.ROOT)}")
