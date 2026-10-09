"""
Proves the workbook's formulas reproduce the pandas pipeline.

openpyxl writes formulas without results, so a spreadsheet engine has to calculate them
first. Two engines are supported:

    # LibreOffice (what CI uses): recalculate a copy, then check the saved values
    python src/recalc_libreoffice.py deliverables/pay_transparency_model.xlsx build/
    python src/check_workbook.py build/pay_transparency_model.xlsx

    # the pure-Python `formulas` package, on a small sample (quick to run locally)
    python src/check_workbook.py --sample 120

The check fails on any formula error anywhere in the workbook, and on any row of the
Reconciliation sheet that does not match.
"""

import argparse
import sys
from pathlib import Path

import excel_twin as et
from excel_twin import ERRORS, values_from_formulas, values_from_saved  # noqa: F401  (used by the tests)
from excel_twin.reconcile import Layout

LAYOUT = Layout(item_col="C", summary_cell="C3")


def check(values: dict, sheet_names: list[str], n_checks: int) -> int:
    """The excel-twin check with this repository's layout; 0 when everything matches."""
    report = et.check(values, n_checks, LAYOUT)
    report.print()
    return 0 if report.ok else 1


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", nargs="?", type=Path,
                    help="a workbook already recalculated by LibreOffice or Excel")
    ap.add_argument("--sample", type=int,
                    help="build a sample workbook of this many workers and evaluate it with `formulas`")
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--formulas", action="store_true",
                    help="calculate the given workbook with `formulas` (slow at full size)")
    args = ap.parse_args()

    if args.sample:
        import pandas as pd

        import build_workbook
        from config import PAYROLL_EXTRACT, ROOT
        from workers import derive
        extract = pd.read_csv(PAYROLL_EXTRACT)
        sample = pd.concat([g.sample(n=min(len(g), args.sample // 4), random_state=args.seed)
                            for _, g in extract.groupby("country")]).sort_values("worker_id")
        path = ROOT / "build" / "sample_model.xlsx"
        model = build_workbook.build(derive(sample), path)
        print(f"sample workbook: {len(sample)} workers, {len(model.checks)} checks")
        values = values_from_formulas(path)
        return check(values, model.wb.sheetnames, len(model.checks))

    if not args.workbook:
        ap.error("give a recalculated workbook, or --sample N")
    from openpyxl import load_workbook
    names = load_workbook(args.workbook, read_only=True).sheetnames
    values = values_from_formulas(args.workbook) if args.formulas else values_from_saved(args.workbook)
    n = sum(1 for (s, r) in values if s.upper() == "RECONCILIATION" and r.startswith("C")
            and r[1:].isdigit() and int(r[1:]) >= 7)
    return check(values, names, n)


if __name__ == "__main__":
    sys.exit(main())
