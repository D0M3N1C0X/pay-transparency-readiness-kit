"""
Builds deliverables/board_briefing.pptx: the ten-slide read-out a board would get.

Charts are native PowerPoint charts, so every number can be inspected and restyled in the
deck. Figures come from the same functions as the report and the workbook.
"""
from datetime import datetime

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

import build_report
from config import DELIVERABLES, ENTITIES, SNAPSHOT
from deterministic import normalise
from workers import load_workers

OUTPUT = DELIVERABLES / "board_briefing.pptx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
NAVY_SOFT = RGBColor(0xCA, 0xD6, 0xE6)
INK = RGBColor(0x1B, 0x24, 0x30)
MUTED = RGBColor(0x5F, 0x6B, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TINT = RGBColor(0xF1, 0xF4, 0xF8)
BLUE = RGBColor(0x2A, 0x78, 0xD6)      # series slot 1
ORANGE = RGBColor(0xEB, 0x68, 0x34)    # series slot 2 and accent
STATUS = {  # reserved status colours, always paired with a text label
    "good": RGBColor(0x0C, 0xA3, 0x0C), "warning": RGBColor(0xFA, 0xB2, 0x19),
    "serious": RGBColor(0xEC, 0x83, 0x5A), "critical": RGBColor(0xD0, 0x3B, 0x3B),
}
TRANSPOSITION_STATUS = {"Transposed": "good", "Partly transposed": "warning", "Draft": "warning",
                        "Not transposed": "critical"}
READINESS_STATUS = {"Ready": "good", "Partial": "warning", "Gap": "critical"}

HEAD, BODY = "Cambria", "Calibri"
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)


def pct(x, d=1):
    return f"{x * 100:.{d}f}%"


def text(slide, x, y, w, h, content, size=14, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False):
    """A text box; `content` is a string or a list of (text, overrides) paragraphs."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    paragraphs = content if isinstance(content, list) else [(content, {})]
    for i, (line, over) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = over.get("align", align)
        p.space_after = Pt(over.get("space", 6))
        run = p.add_run()
        run.text = line
        f = run.font
        f.name = over.get("font", font)
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.italic = over.get("italic", italic)
        f.color.rgb = over.get("color", color)
    return box


def rect(slide, x, y, w, h, fill, rounded=True, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if rounded:
        shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def chip(slide, x, y, label, fill, color=WHITE, w=None):
    """The deck's motif: a small rounded label naming the article a slide rests on."""
    width = w or Inches(0.12 * len(label) + 0.35)
    shape = rect(slide, x, y, width, Inches(0.34), fill)
    shape.adjustments[0] = 0.5
    tf = shape.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.name, run.font.size, run.font.bold, run.font.color.rgb = BODY, Pt(11), True, color
    return shape


def status_chip(slide, x, y, label, level, w=Inches(1.7)):
    dark_text = level in ("warning", "serious")
    return chip(slide, x, y, label, STATUS[level], color=INK if dark_text else WHITE, w=w)


def heading(slide, kicker, title, subtitle=None):
    chip(slide, MARGIN, Inches(0.45), kicker, NAVY)
    text(slide, MARGIN, Inches(0.95), W - 2 * MARGIN, Inches(0.9), title, size=30, bold=True, font=HEAD)
    if subtitle:
        text(slide, MARGIN, Inches(1.75), W - 2 * MARGIN, Inches(0.5), subtitle, size=14, color=MUTED)


def footer(slide, n):
    text(slide, MARGIN, H - Inches(0.45), Inches(9), Inches(0.3),
         "Pay transparency readiness · dry run on synthetic data · not legal advice", size=10, color=MUTED)
    text(slide, W - MARGIN - Inches(1), H - Inches(0.45), Inches(1), Inches(0.3), str(n), size=10,
         color=MUTED, align=PP_ALIGN.RIGHT)


def fit(values, width):
    """One font size for a row of big numbers, small enough for the longest to stay on one line."""
    longest = max(len(v) for v in values)
    return max(20, min(40, int(width / 12700 / longest / 0.68)))   # EMU -> pt; ~0.68 em per character


def stat_row(slide, y, items, gap=Inches(0.3)):
    """items: [(value, label, colour)] laid out as equal cards across the slide."""
    w = (W - 2 * MARGIN - gap * (len(items) - 1)) / len(items)
    pad = Inches(0.25)
    size = fit([v for v, _, _ in items], w - 2 * pad)
    for i, (value, label, colour) in enumerate(items):
        x = MARGIN + i * (w + gap)
        rect(slide, x, y, w, Inches(2.0), TINT)
        text(slide, x + pad, y + Inches(0.25), w - 2 * pad, Inches(0.8), value, size=size, bold=True,
             color=colour, font=HEAD, anchor=MSO_ANCHOR.BOTTOM)
        text(slide, x + pad, y + Inches(1.15), w - 2 * pad, Inches(0.8), label, size=14, color=INK)


def style_chart(chart, colours, fmt="0.0%", legend=True):
    chart.font.name = BODY
    chart.font.size = Pt(12)
    chart.font.color.rgb = INK
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = fmt
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(11)
    labels.font.color.rgb = INK
    for series, colour in zip(plot.series, colours):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = colour
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = False
    value_axis.visible = False
    value_axis.minimum_scale = 0
    cat = chart.category_axis
    cat.format.line.color.rgb = RGBColor(0xC3, 0xC2, 0xB7)
    cat.tick_labels.font.size = Pt(12)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)


def build(w=None):
    w = load_workers() if w is None else w
    f = build_report.facts(w)
    codes, ind, cats, obl, ready, flagged = (f[k] for k in ("codes", "ind", "cats", "obl", "ready", "flagged"))
    G = ind["Group"]
    first = obl["first_report"].min()
    remedy_by = obl["remedy_by"].min()
    counts = ready["status"].value_counts()
    status = f["obl"].set_index("country")["transposition"]
    day = build_report.day

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    def new_slide(dark=False):
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = NAVY if dark else WHITE
        return s

    # 1 - title ----------------------------------------------------------------------------
    s = new_slide(dark=True)
    chip(s, MARGIN, Inches(1.2), "Directive (EU) 2023/970", ORANGE)
    text(s, MARGIN, Inches(1.9), Inches(10.5), Inches(2.2),
         "Pay transparency: where we stand before the first report", size=44, bold=True, color=WHITE, font=HEAD)
    text(s, MARGIN, Inches(4.2), Inches(10), Inches(0.9),
         f"Dry run on pay at {day(SNAPSHOT)} · four employers · {G['workers']:,.0f} workers",
         size=20, color=NAVY_SOFT)
    text(s, MARGIN, Inches(6.2), Inches(10), Inches(0.8),
         [("Board briefing · September 2026", {"bold": True, "color": WHITE}),
          ("Synthetic data built for demonstration. Not legal advice.", {"color": NAVY_SOFT, "size": 12})],
         size=14)
    s.notes_slide.notes_text_frame.text = (
        "Purpose: tell the board where the organisation stands against the EU Pay Transparency Directive "
        "before the first statutory report, and what decisions are needed now.")

    # 2 - the answer -------------------------------------------------------------------------
    s = new_slide()
    heading(s, "The answer", "Four employers report every year from June 2027, and none is ready yet")
    stat_row(s, Inches(2.3), [
        (f"{first:%b %Y}", f"first report, due {day(first)}, for all four employers; then every year (Art. 9(2))",
         NAVY),
        (pct(G["a_mean_gap_hourly"]), "gender pay gap on hourly pay across the group (Art. 9(1)(a))", ORANGE),
        (f"{len(flagged)} of {len(cats)}", "categories of workers at 5% or more: justify or remedy (Art. 10)", ORANGE),
        (f"€{flagged['remedy_cost'].sum() / 1e6:.1f}M", "a year, at most, to close those gaps outright", NAVY),
    ])
    text(s, MARGIN, Inches(4.8), W - 2 * MARGIN, Inches(1.6), [
        (f"Readiness: {counts.get('Ready', 0)} of {len(ready)} obligations in place, {counts.get('Partial', 0)} partly, "
         f"{counts.get('Gap', 0)} missing.", {"bold": True, "size": 18}),
        (f"Without a justification or a remedy by {day(remedy_by)}, each flagged category triggers a joint pay "
         "assessment with workers' representatives - and, where the transparency duties were not met, the "
         "burden of proof in a pay claim shifts to the employer (Art. 18(2)).", {"size": 16, "color": MUTED}),
    ])
    footer(s, 2)
    s.notes_slide.notes_text_frame.text = (
        "Lead with the four numbers. The cost is an upper bound: it lifts the lower-paid sex's average hourly "
        "pay in every flagged category, before any part of the gap is justified.")

    # 3 - what the directive asks ------------------------------------------------------------
    s = new_slide()
    heading(s, "Articles 9, 10 and 34", "The timetable is fixed; the national rules are not")
    steps = [
        ("7 Jun 2026", "Transposition deadline", "Member States had to bring the Directive into national law."),
        (day(first, "%b"), "First pay report", "250+ workers: every year. 150-249: every three years."),
        (day(remedy_by, "%b"), "Justify or remedy", "Six months to explain or close any category gap of 5% or more."),
        ("Then", "Joint pay assessment", "With workers' representatives, if a gap is neither justified nor remedied."),
    ]
    sw = (W - 2 * MARGIN - Inches(0.9)) / 4
    for i, (when, what, detail) in enumerate(steps):
        x = MARGIN + i * (sw + Inches(0.3))
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, x, Inches(2.35), Inches(0.28), Inches(0.28))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ORANGE if i else NAVY
        dot.line.fill.background()
        if i < len(steps) - 1:
            line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x + Inches(0.34), Inches(2.49), x + sw + Inches(0.24), Inches(2.49))
            line.line.color.rgb = RGBColor(0xC3, 0xC2, 0xB7)
            line.line.width = Pt(1.5)
        text(s, x, Inches(2.8), sw, Inches(0.45), when, size=20, bold=True, font=HEAD, color=NAVY)
        text(s, x, Inches(3.25), sw, Inches(0.4), what, size=15, bold=True)
        text(s, x, Inches(3.65), sw, Inches(0.8), detail, size=13, color=MUTED)
    y = Inches(4.75)
    text(s, MARGIN, y, Inches(8), Inches(0.4), "Status of national law, 16 September 2026", size=14, bold=True)
    notes = {
        "IT": "D.Lgs. 96/2026, in force 7 June 2026; collective agreements anchor the categories",
        "PL": "Recruitment rules in force; the reporting act is a draft",
        "DE": "No act yet; the 2017 pay transparency act still applies",
        "ES": "Draft Royal Decree, consulted in August 2026",
    }
    for i, code in enumerate(codes):
        col, row = i % 2, i // 2
        x = MARGIN + col * Inches(6.2)
        yy = y + Inches(0.5) + row * Inches(0.75)
        text(s, x, yy + Inches(0.03), Inches(1.6), Inches(0.35), ENTITIES[code], size=14, bold=True)
        status_chip(s, x + Inches(1.65), yy, status[code], TRANSPOSITION_STATUS[status[code]])
        text(s, x + Inches(3.5), yy - Inches(0.02), Inches(2.6), Inches(0.7), notes[code], size=11, color=MUTED)
    footer(s, 3)
    s.notes_slide.notes_text_frame.text = (
        "Only Italy has transposed in full among our four countries. The Directive's dates still frame the "
        "plan; national acts can add detail on method and filing, which is why the model keeps every choice "
        "in a Settings sheet.")

    # 4 - headline gaps --------------------------------------------------------------------
    s = new_slide()
    heading(s, "Article 9(1)(a) and (c)", "Women earn less per hour in every employer",
            "Gender pay gap on gross hourly pay, base plus bonus, commission and allowances")
    data = CategoryChartData()
    data.categories = [ENTITIES[c] for c in codes]
    data.add_series("Mean gap", [float(ind.loc["a_mean_gap_hourly", c]) for c in codes])
    data.add_series("Median gap", [float(ind.loc["c_median_gap_hourly", c]) for c in codes])
    frame = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(2.4), Inches(7.6), Inches(4.4), data)
    style_chart(frame.chart, [BLUE, ORANGE])
    x = Inches(8.7)
    text(s, x, Inches(2.6), Inches(4), Inches(4), [
        ("Median below mean everywhere", {"bold": True, "size": 17}),
        ("The gap is concentrated at the top of the pay distribution.", {"color": MUTED, "space": 16}),
        (f"Base pay alone: {pct(G['base_mean_gap_hourly'])}", {"bold": True, "size": 17}),
        (f"Components add {(G['a_mean_gap_hourly'] - G['base_mean_gap_hourly']) * 100:.1f} points to the group "
         "gap - the part a base-pay analysis misses.", {"color": MUTED, "space": 16}),
        (f"Group: {pct(G['a_mean_gap_hourly'])} mean, {pct(G['c_median_gap_hourly'])} median", {"bold": True, "size": 17}),
        ("A management view; each employer reports on its own.", {"color": MUTED}),
    ], size=14)
    footer(s, 4)
    s.notes_slide.notes_text_frame.text = (
        "The 15.4% base-pay figure is the same number the earlier people-analytics review found, which is a "
        "useful cross-check: the two analyses agree before components are added.")

    # 5 - quartile bands -----------------------------------------------------------------------
    s = new_slide()
    heading(s, "Article 9(1)(f)", "Women fill the lowest pay band, not the highest",
            "Share of women in the lowest and highest of four equal pay bands, by employer")
    data = CategoryChartData()
    data.categories = [ENTITIES[c] for c in codes]
    data.add_series("Lowest band", [float(ind.loc["f_q1_share_women", c]) for c in codes])
    data.add_series("Highest band", [float(ind.loc["f_q4_share_women", c]) for c in codes])
    frame = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MARGIN, Inches(2.4), Inches(7.6), Inches(4.4), data)
    style_chart(frame.chart, [ORANGE, BLUE], fmt="0%")
    frame.chart.category_axis.reverse_order = True
    share = (w["gender"] == "F").groupby(w["job_level"]).mean()
    text(s, Inches(8.7), Inches(2.6), Inches(4), Inches(4), [
        (f"{pct(G['f_q1_share_women'], 0)} of the lowest band, {pct(G['f_q4_share_women'], 0)} of the highest",
         {"bold": True, "size": 17}),
        ("Across the group, against 50% parity.", {"color": MUTED, "space": 16}),
        ("Vertical segregation is the main driver", {"bold": True, "size": 17}),
        (f"Women are {pct(share['L1'], 0)} of entry roles and {pct(share['L6'], 0)} of country heads of function. "
         "Promotion and hiring into senior roles move this number more than pay adjustments do.",
         {"color": MUTED}),
    ], size=14)
    footer(s, 5)
    s.notes_slide.notes_text_frame.text = "Quartile bands are four equal groups of workers ranked by hourly pay within each employer."

    # 6 - components ----------------------------------------------------------------------------
    s = new_slide()
    heading(s, "Article 9(1)(b), (d) and (e)", "Bonus, commission and allowances widen the gap")
    rate = f["sales_rate"]
    fewer = (G["e_share_men_receiving"] - G["e_share_women_receiving"]) * 100
    stat_row(s, Inches(2.1), [
        (pct(G["b_mean_gap_components"]), "gap in average components, among workers who receive any", ORANGE),
        (f"{fewer:.1f} pts", f"fewer women receive any component: {pct(G['e_share_women_receiving'], 0)} "
                             f"against {pct(G['e_share_men_receiving'], 0)} of men", NAVY),
        (f"{(rate['M'] - rate['F']) * 100:.1f} pts", f"lower commission rate for women in Sales: {pct(rate['F'])} "
                                                   f"of base pay against {pct(rate['M'])}", ORANGE),
    ])
    rect(s, MARGIN, Inches(4.5), W - 2 * MARGIN, Inches(2.2), TINT)
    text(s, MARGIN + Inches(0.35), Inches(4.75), Inches(5.6), Inches(1.8), [
        ("Eligibility follows seniority", {"bold": True, "size": 17}),
        ("Bonus starts at L3 and car allowances at L5, where fewer women work. Structural, and visible in the "
         "quartile bands.", {"color": MUTED}),
    ], size=14)
    text(s, MARGIN + Inches(6.3), Inches(4.75), Inches(5.6), Inches(1.8), [
        ("Commission follows territory", {"bold": True, "size": 17}),
        ("The commission rules are the same for everyone; women are given territories with lower potential. "
         "Neutral on paper, unequal in effect: the kind of rule a joint pay assessment exists to find.",
         {"color": MUTED}),
    ], size=14)
    footer(s, 6)
    s.notes_slide.notes_text_frame.text = (
        "The territory effect is planted in the synthetic data on purpose, and documented, so the screen can be "
        "shown finding it. In a real client this is a question for Sales operations: how are territories assigned?")

    # 7 - Article 10 screen ------------------------------------------------------------------------
    s = new_slide()
    heading(s, "Article 10", f"{len(flagged)} categories to justify or remedy by {day(remedy_by, '%b')}")
    top = flagged.sort_values("remedy_cost", ascending=False)
    rows, cols = len(top) + 1, 6
    table = s.shapes.add_table(rows, cols, MARGIN, Inches(1.95), W - 2 * MARGIN, Inches(0.36) * rows).table
    widths = [2.0, 1.1, 1.4, 1.3, 4.6, 1.73]
    for i, wd in enumerate(widths):
        table.columns[i].width = Inches(wd)
    heads = ["Employer", "Category", "Women / men", "Hourly gap", "Main driver", "Cost to close"]
    for j, h_ in enumerate(heads):
        c = table.cell(0, j)
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        c.text = h_
    for i, r in enumerate(top.itertuples(), start=1):
        values = [ENTITIES[r.country], r.category.replace("Cat ", ""), f"{r.women} / {r.men}", pct(r.gap_hourly),
                  build_report.driver(r._asdict()), f"€{r.remedy_cost / 1000:,.0f}k"]
        for j, v in enumerate(values):
            c = table.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = TINT if i % 2 else WHITE
            c.text = v
    for i in range(rows):
        table.rows[i].height = Inches(0.36)
        for j in range(cols):
            c = table.cell(i, j)
            c.margin_top = c.margin_bottom = Emu(30000)
            for p in c.text_frame.paragraphs:
                p.alignment = PP_ALIGN.RIGHT if j in (2, 3, 5) else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name, run.font.size = BODY, Pt(12)
                    run.font.bold = i == 0
                    run.font.color.rgb = WHITE if i == 0 else INK
    footer(s, 7)
    s.notes_slide.notes_text_frame.text = (
        "Flag rule: a difference of 5% or more in either direction, on hourly or annual pay. Categories come from "
        "a gender-neutral job evaluation on skills, effort, responsibility and working conditions. Two "
        "categories have fewer than five workers of one sex; their figures go only to workers' representatives, "
        "the labour inspectorate or the equality body.")

    # 8 - readiness --------------------------------------------------------------------------------
    s = new_slide()
    heading(s, "Articles 4 to 12", "Most of what workers will see is not built yet")
    for i, level in enumerate(["Ready", "Partial", "Gap"]):
        x = MARGIN + i * Inches(1.75)
        text(s, x, Inches(2.05), Inches(1.5), Inches(0.8), str(counts.get(level, 0)), size=40, bold=True, font=HEAD)
        status_chip(s, x, Inches(2.9), level, READINESS_STATUS[level], w=Inches(1.3))
    gaps = ready[ready["status"] == "Gap"]
    text(s, MARGIN + Inches(5.6), Inches(2.2), Inches(6.5), Inches(1.0),
         "The missing pieces are the ones workers and applicants see first: pay ranges, published criteria, "
         "an answer to their questions.", size=16, color=MUTED)
    cols = 4
    gw = (W - 2 * MARGIN - Inches(0.25) * (cols - 1)) / cols
    for i, r in enumerate(gaps.itertuples()):
        col, row = i % cols, i // cols
        x = MARGIN + col * (gw + Inches(0.25))
        y = Inches(3.55) + row * Inches(1.7)
        rect(s, x, y, gw, Inches(1.5), TINT)
        text(s, x + Inches(0.18), y + Inches(0.14), gw - Inches(0.36), Inches(1.25), [
            (f"{r.article} · {r.owner}", {"bold": True, "size": 11, "color": MUTED, "space": 3}),
            (r.requirement, {"size": 11}),
        ])
    footer(s, 8)
    s.notes_slide.notes_text_frame.text = "Statuses are illustrative for the demonstration organisation."

    # 9 - next 90 days -----------------------------------------------------------------------------
    s = new_slide()
    heading(s, "The plan", "The next 90 days")
    plan = [
        ("Agree job evaluation", "With workers' representatives, starting from the Italian collective agreement."),
        ("Justify or budget flags", "Most expensive first: a documented, gender-neutral reason or a remedy."),
        ("Review Sales territories", "And every other rule that decides who can earn a component."),
        ("Fix recruitment now", "Pay ranges in vacancies, no pay-history questions, neutral titles."),
        ("Open information route", "Written answers within two months; an annual notice to workers."),
        ("Rerun on 2026 payroll", "Once national rules are final; management sign-off after consultation."),
    ]
    pw = (W - 2 * MARGIN - Inches(0.6)) / 3
    for i, (what, detail) in enumerate(plan):
        col, row = i % 3, i // 3
        x = MARGIN + col * (pw + Inches(0.3))
        y = Inches(2.1) + row * Inches(2.35)
        rect(s, x, y, pw, Inches(2.05), TINT)
        num = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.3), y + Inches(0.3), Inches(0.5), Inches(0.5))
        num.fill.solid()
        num.fill.fore_color.rgb = NAVY
        num.line.fill.background()
        tf = num.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(i + 1)
        run.font.name, run.font.size, run.font.bold, run.font.color.rgb = BODY, Pt(16), True, WHITE
        text(s, x + Inches(0.3), y + Inches(0.9), pw - Inches(0.6), Inches(0.4), what, size=16, bold=True)
        text(s, x + Inches(0.3), y + Inches(1.32), pw - Inches(0.6), Inches(0.7), detail, size=13, color=MUTED)
    footer(s, 9)

    # 10 - decisions -------------------------------------------------------------------------------
    s = new_slide(dark=True)
    chip(s, MARGIN, Inches(0.6), "For decision today", ORANGE)
    text(s, MARGIN, Inches(1.15), Inches(12), Inches(0.9), "Three decisions for the board", size=36, bold=True,
         color=WHITE, font=HEAD)
    decisions = [
        ("Mandate the job evaluation", "Owner: Reward, with workers' representatives in each country. Categories "
                                       "decide every later number."),
        (f"Reserve up to €{flagged['remedy_cost'].sum() / 1e6:.1f}M a year",
         "A ceiling for remedies, released category by category once justifications are reviewed."),
        ("Name owners for Articles 5 to 7", "Recruitment, pay criteria and information requests, live before the "
                                           "first report."),
    ]
    dw = (W - 2 * MARGIN - Inches(0.6)) / 3
    for i, (what, detail) in enumerate(decisions):
        x = MARGIN + i * (dw + Inches(0.3))
        box = rect(s, x, Inches(2.5), dw, Inches(3.1), RGBColor(0x2A, 0x4B, 0x74))
        text(s, x + Inches(0.35), Inches(2.8), dw - Inches(0.7), Inches(0.6), str(i + 1), size=36, bold=True,
             color=ORANGE, font=HEAD)
        text(s, x + Inches(0.35), Inches(3.6), dw - Inches(0.7), Inches(0.9), what, size=20, bold=True, color=WHITE)
        text(s, x + Inches(0.35), Inches(4.45), dw - Inches(0.7), Inches(1.1), detail, size=14, color=NAVY_SOFT)
    text(s, MARGIN, Inches(6.3), Inches(12), Inches(0.8), [
        ("Method: every figure computed twice - in Python and in a live Excel model - and reconciled. Legal "
         "status checked on 16 September 2026, sources in docs/verification.md. Synthetic data; not legal advice.",
         {"color": NAVY_SOFT, "size": 11}),
        ("github.com/D0M3N1C0X/pay-transparency-readiness-kit", {"color": WHITE, "size": 11, "bold": True}),
    ])
    s.notes_slide.notes_text_frame.text = (
        "Ask for the three decisions explicitly. The budget is a ceiling, not a spend: justified gaps need "
        "documentation, not money.")

    props = prs.core_properties
    stamp = datetime(2026, 1, 1)
    props.author = props.last_modified_by = "Domenico Perroni"
    props.title = "Pay transparency: where we stand before the first report"
    props.created = props.modified = stamp
    props.revision = 1
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    normalise(OUTPUT)
    return OUTPUT


def main() -> None:
    path = build()
    print(f"deck -> {path.relative_to(path.parents[1])}")


if __name__ == "__main__":
    main()
