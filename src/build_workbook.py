"""
Builds deliverables/pay_transparency_model.xlsx, the client-facing model.

Every figure in the workbook is a live formula over the Payroll sheet: change a salary, a
job evaluation score or a threshold and the report recalculates. The Reconciliation sheet
holds the values this pipeline computed in pandas and checks each one against the formula
that should reproduce it, so the model and the code vouch for each other.

Conventions: blue text = input, black = formula, green = link to another sheet, yellow fill =
key assumption. Arial throughout. Only functions that Excel, LibreOffice and Numbers all
evaluate (no dynamic arrays).
"""
import math
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.hyperlink import Hyperlink

import article9
import article10
import obligations
from build_extract import COLUMNS as INPUT_COLUMNS
from config import (ASSESSMENT_DATE, BANDS, DELIVERABLES, ENTITIES, FACTOR_WEIGHTS, FACTORS,
                    FULL_TIME_WEEKLY_HOURS, JPA_THRESHOLD, MIN_CELL, REMEDY_MONTHS, SIZE_BANDS,
                    SNAPSHOT, WEEKS_PER_YEAR)
from deterministic import normalise
from workers import load_job_evaluation, load_workers

OUTPUT = DELIVERABLES / "pay_transparency_model.xlsx"
REPO = "https://github.com/D0M3N1C0X/pay-transparency-readiness-kit"

# ---- style ------------------------------------------------------------------------------
FONT = "Arial"
BLUE, GREEN, INK, MUTED, WHITE = "0000FF", "008000", "1B2430", "5F6B7A", "FFFFFF"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
BAND_FILL = PatternFill("solid", fgColor="EEF2F7")
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="C9D1DC")
BOX = Border(bottom=THIN)

PCT = "0.0%"
COUNT = "#,##0"
EUR = '"€"#,##0;-"€"#,##0;"-"'
EUR_H = '"€"0.00'
DATE = "d mmm yyyy"

STATUS_FILLS = {
    "Justify or remedy": "F8D7B0", "Below 5%": "D5ECDC", "No comparison: one sex only": "E4E7EC",
    "Ready": "D5ECDC", "Partial": "FBE8B5", "Gap": "F5C6C2", "Yes": "D5ECDC", "No": "F5C6C2",
}


def font(color=INK, bold=False, italic=False, size=10):
    return Font(name=FONT, color=color, bold=bold, italic=italic, size=size)


def put(ws, ref, value, *, color=INK, bold=False, italic=False, size=10, fmt=None, fill=None,
        wrap=False, align=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font(color, bold, italic, size)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if wrap or align:
        cell.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top")
    return cell


def array(ws, ref, formula, fmt=None):
    ws[ref] = ArrayFormula(ref, formula)
    ws[ref].font = font()
    if fmt:
        ws[ref].number_format = fmt


def header(ws, row, labels, start_col=1, height=30):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=label)
        c.font = font(WHITE, bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height


def title(ws, text, subtitle=None):
    put(ws, "A1", text, bold=True, size=14)
    if subtitle:
        put(ws, "A2", subtitle, color=MUTED, italic=True)


def widths(ws, spec: dict):
    for col, width in spec.items():
        ws.column_dimensions[col].width = width


def status_colours(ws, rng):
    for text, colour in STATUS_FILLS.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{text}"'],
                                                      fill=PatternFill("solid", fgColor=colour)))


def blank(x):
    """pandas NaN -> empty cell, so the reconciliation compares like with like."""
    if x is None:
        return None
    if isinstance(x, float) and math.isnan(x):
        return None
    return x


class Model:
    def __init__(self, workers: pd.DataFrame, grid: pd.DataFrame):
        self.w = workers
        self.grid = grid
        self.countries = list(ENTITIES)
        self.wb = Workbook()
        self.refs = {}          # named anchors: key -> "Sheet!$C$5"
        self.checks = []        # (area, entity, item, python value, workbook ref)

    # -- helpers --------------------------------------------------------------------------
    def col(self, name):
        """Absolute range of a Payroll column."""
        letter = self.payroll_cols[name]
        return f"Payroll!${letter}$2:${letter}${self.last}"

    def check(self, area, entity, item, value, ref):
        self.checks.append((area, entity, item, blank(value), ref))

    # -- sheets ---------------------------------------------------------------------------
    def build(self, path: Path):
        wb = self.wb
        names = ["Cover", "Summary", "Article 9", "Categories", "Obligations", "Readiness",
                 "Job Evaluation", "Payroll", "Settings", "Reconciliation"]
        wb.active.title = names[0]
        for n in names[1:]:
            wb.create_sheet(n)
        # Inputs first, so the output sheets can point at them.
        self.settings(wb["Settings"])
        self.job_evaluation(wb["Job Evaluation"])
        self.payroll(wb["Payroll"])
        self.article9(wb["Article 9"])
        self.categories(wb["Categories"])
        self.obligations(wb["Obligations"])
        self.readiness(wb["Readiness"])
        self.summary(wb["Summary"])
        self.reconciliation(wb["Reconciliation"])
        self.cover(wb["Cover"], names)

        for ws in wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb.calculation.fullCalcOnLoad = True
        stamp = datetime(2026, 1, 1)
        wb.properties.creator = "Domenico Perroni"
        wb.properties.title = "Pay Transparency Readiness Model"
        wb.properties.created = stamp
        wb.properties.modified = stamp
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        normalise(path)

    def settings(self, ws):
        title(ws, "Settings", "Every assumption the model uses. Blue cells can be changed; "
                               "the yellow ones move the results most.")
        widths(ws, {"A": 3, "B": 46, "C": 16, "D": 22, "E": 16, "F": 70})
        r = 4
        put(ws, f"B{r}", "General", bold=True, size=11)
        rows = [
            ("snapshot", "Snapshot date (population: workers active on this date)", SNAPSHOT, DATE, False,
             "Dry run on annualised pay of workers active at the snapshot. The statutory report covers pay in the previous calendar year."),
            ("assessment", "Date of this assessment and of the legal status notes", ASSESSMENT_DATE, DATE, False, ""),
            ("weeks", "Weeks per year (annual hours = weekly hours x weeks)", WEEKS_PER_YEAR, "0", True,
             "Hourly pay = annual pay / annual contractual hours."),
            ("ft_hours", "Contractual full-time week, hours (all entities)", FULL_TIME_WEEKLY_HOURS, "0", False,
             "Synthetic assumption used when the extract was built; weekly hours per worker are in Payroll."),
            ("threshold", "Joint pay assessment threshold (Article 10(1)(a))", JPA_THRESHOLD, PCT, True,
             "A difference of at least 5% in any category, in either direction, on hourly or annual pay."),
            ("months", "Months to justify or remedy after submission (Article 10(1)(c))", REMEDY_MONTHS, "0", False, ""),
            ("min_cell", "Fewest workers of one sex before a category average is shared widely", MIN_CELL, "0", True,
             "Article 12(3) lets Member States restrict disclosure that would reveal an identifiable worker's pay. The Directive sets no number; 5 is this kit's choice."),
        ]
        r = 5
        for key, label, value, fmt, key_assumption, note in rows:
            put(ws, f"B{r}", label)
            put(ws, f"C{r}", value, color=BLUE, fmt=fmt, fill=YELLOW if key_assumption else None)
            put(ws, f"F{r}", note, color=MUTED, italic=True, wrap=True)
            self.refs[key] = f"Settings!$C${r}"
            r += 1

        r += 1
        put(ws, f"B{r}", "Job evaluation weights (Article 4(4) criteria), points per score point", bold=True, size=11)
        r += 1
        first = r
        for f in FACTORS:
            put(ws, f"B{r}", f.replace("_", " ").capitalize())
            put(ws, f"C{r}", FACTOR_WEIGHTS[f], color=BLUE, fmt="0", fill=YELLOW)
            self.refs[f"weight_{f}"] = f"Settings!$C${r}"
            r += 1
        put(ws, f"B{r}", "Total (must be 100)", bold=True)
        put(ws, f"C{r}", f"=SUM(C{first}:C{r - 1})", bold=True, fmt="0")
        put(ws, f"F{r}", "With scores from 1 to 5, a role scores between 100 and 500 points.", color=MUTED, italic=True)

        r += 2
        put(ws, f"B{r}", "Categories of workers: point bands (lower limit inclusive)", bold=True, size=11)
        r += 1
        header(ws, r, ["Lower limit, points", "Category"], start_col=2, height=18)
        r += 1
        first = r
        for floor, name in BANDS:
            put(ws, f"B{r}", floor, color=BLUE, fmt="0", align="left")
            put(ws, f"C{r}", name, color=BLUE)
            r += 1
        self.refs["band_floor"] = f"Settings!$B${first}:$B${r - 1}"
        self.refs["band_name"] = f"Settings!$C${first}:$C${r - 1}"

        r += 1
        put(ws, f"B{r}", "Reporting duty by size (Article 9(2)-(4)), ascending", bold=True, size=11)
        r += 1
        header(ws, r, ["Minimum workers", "Band", "Frequency", "First report"], start_col=2, height=18)
        r += 1
        first = r
        for floor, label, frequency, first_report in sorted(SIZE_BANDS):
            put(ws, f"B{r}", floor, color=BLUE, fmt="0", align="left")
            put(ws, f"C{r}", label, color=BLUE)
            put(ws, f"D{r}", frequency, color=BLUE)
            put(ws, f"E{r}", first_report if first_report else "n/a", color=BLUE, fmt=DATE)
            r += 1
        for key, letter in (("size_floor", "B"), ("size_label", "C"), ("size_freq", "D"), ("size_first", "E")):
            self.refs[key] = f"Settings!${letter}${first}:${letter}${r - 1}"
        put(ws, f"F{first}", "Member States may extend duties to employers under 100 (Article 9(5)).",
            color=MUTED, italic=True, wrap=True)

        r += 1
        put(ws, f"B{r}", "Employers (one legal entity per Member State)", bold=True, size=11)
        r += 1
        header(ws, r, ["Code", "Entity"], start_col=2, height=18)
        r += 1
        for code, name in ENTITIES.items():
            put(ws, f"B{r}", code, color=BLUE)
            put(ws, f"C{r}", name, color=BLUE)
            r += 1

    def job_evaluation(self, ws):
        title(ws, "Job evaluation: gender-neutral scoring of every role",
              "Scores 1-5 on the four criteria of Article 4(4). Agree them with workers' representatives; "
              "roles with similar points form one category of workers, whatever the department.")
        widths(ws, {"A": 24, "B": 18, "C": 7, "D": 9, "E": 9, "F": 14, "G": 12, "H": 9, "I": 10, "J": 90})
        header(ws, 4, ["Role", "Department", "Level", "Skills", "Effort", "Responsibility",
                       "Working conditions", "Points", "Category", "Rationale"])
        weights = [self.refs[f"weight_{f}"] for f in FACTORS]
        for i, g in enumerate(self.grid.itertuples(index=False), start=5):
            put(ws, f"A{i}", g.role, color=BLUE)
            put(ws, f"B{i}", g.department, color=BLUE)
            put(ws, f"C{i}", g.job_level, color=BLUE)
            for letter, f in zip("DEFG", FACTORS):
                put(ws, f"{letter}{i}", int(getattr(g, f)), color=BLUE, fmt="0", align="center")
            put(ws, f"H{i}", "=" + "+".join(f"{l}{i}*{w}" for l, w in zip("DEFG", weights)), fmt="0")
            put(ws, f"I{i}", f"=INDEX({self.refs['band_name']},MATCH(H{i},{self.refs['band_floor']},1))")
            put(ws, f"J{i}", g.rationale, color=BLUE, wrap=False)
            self.check("Job evaluation", "", f"{g.role}: points", int(g.points), f"'Job Evaluation'!H{i}")
            self.check("Job evaluation", "", f"{g.role}: category", g.category, f"'Job Evaluation'!I{i}")
        last = 4 + len(self.grid)
        self.refs["je_role"] = f"'Job Evaluation'!$A$5:$A${last}"
        self.refs["je_category"] = f"'Job Evaluation'!$I$5:$I${last}"
        ws.freeze_panes = "B5"
        ws.auto_filter.ref = f"A4:J{last}"

    def payroll(self, ws):
        w = self.w
        n = len(w)
        self.last = n + 1
        derived = ["annual_hours", "components", "total_pay", "hourly_base", "hourly_components",
                   "hourly_total", "receives_components", "category", "rank_in_entity", "quartile"]
        names = INPUT_COLUMNS + derived
        self.payroll_cols = {name: get_column_letter(i) for i, name in enumerate(names, start=1)}
        c = self.payroll_cols
        header(ws, 1, names, height=32)
        for name in derived:
            ws[f"{c[name]}1"].fill = PatternFill("solid", fgColor="3E5C76")

        money = {"base_pay", "bonus", "commission", "shift_allowance", "car_allowance"}
        values = w[INPUT_COLUMNS].itertuples(index=False)
        blue = font(BLUE)
        black = font()
        last = self.last
        for r, row in enumerate(values, start=2):
            for name, value in zip(INPUT_COLUMNS, row):
                cell = ws[f"{c[name]}{r}"]
                cell.value = value.item() if hasattr(value, "item") else value
                cell.font = blue
                if name in money:
                    cell.number_format = EUR
            L = {k: f"{v}{r}" for k, v in c.items()}
            formulas = {
                "annual_hours": f"={L['weekly_hours']}*{self.refs['weeks']}",
                "components": f"={L['bonus']}+{L['commission']}+{L['shift_allowance']}+{L['car_allowance']}",
                "total_pay": f"={L['base_pay']}+{L['components']}",
                "hourly_base": f"={L['base_pay']}/{L['annual_hours']}",
                "hourly_components": f"={L['components']}/{L['annual_hours']}",
                "hourly_total": f"={L['total_pay']}/{L['annual_hours']}",
                "receives_components": f"=IF({L['components']}>0,1,0)",
                "category": f"=INDEX({self.refs['je_category']},MATCH({L['role']},{self.refs['je_role']},0))",
                # Rank within the employer by hourly pay, ties broken by row order. SUMPRODUCT
                # compares the numbers themselves; a COUNTIFS "<"&value criterion would round
                # the value to 15 digits first and miscount it against itself.
                "rank_in_entity": (f"=SUMPRODUCT((${c['country']}$2:${c['country']}${last}={L['country']})"
                                   f"*(${c['hourly_total']}$2:${c['hourly_total']}${last}<{L['hourly_total']}))"
                                   f"+SUMPRODUCT((${c['country']}$2:{L['country']}={L['country']})"
                                   f"*(${c['hourly_total']}$2:{L['hourly_total']}={L['hourly_total']}))"),
                "quartile": (f"=MIN(4,INT(({L['rank_in_entity']}-1)*4"
                             f"/COUNTIF(${c['country']}$2:${c['country']}${last},{L['country']}))+1)"),
            }
            fmts = {"annual_hours": COUNT, "components": EUR, "total_pay": EUR, "hourly_base": EUR_H,
                    "hourly_components": EUR_H, "hourly_total": EUR_H}
            for name, formula in formulas.items():
                cell = ws[f"{c[name]}{r}"]
                cell.value = formula
                cell.font = black
                if name in fmts:
                    cell.number_format = fmts[name]
        for name, letter in c.items():
            ws.column_dimensions[letter].width = 13
        for name in ("entity", "department", "role"):
            ws.column_dimensions[c[name]].width = 20
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(names))}{last}"

    def article9(self, ws):
        c = self.col
        codes = self.countries
        cols = {code: get_column_letter(3 + i) for i, code in enumerate(codes)}
        cols["Group"] = get_column_letter(3 + len(codes))
        title(ws, "Article 9(1) reporting items, per employer",
              "Gaps are men's pay level minus women's, as a share of men's (Article 3(1)(c)). "
              "Positive means women are paid less. The group column is a management view, not a statutory item.")
        widths(ws, {"A": 3, "B": 60, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 70})
        header(ws, 4, ["", "Reporting item", *codes, "Group", "How it is calculated"])
        for code in codes:
            put(ws, f"{cols[code]}5", f"=INDEX(Settings!$C$1:$C$200,MATCH(\"{code}\",Settings!$B$1:$B$200,0))",
                color=GREEN, italic=True)
        put(ws, "B5", "Employer", italic=True, color=MUTED)

        # Workings first (rows 30+), so the reporting items can refer to them.
        work = {}
        header(ws, 29, ["", "Workings", *codes, "Group", ""], height=18)
        r = 30
        specs = [
            ("women", "Women", "count"), ("men", "Men", "count"), ("workers", "Workers", "count"),
            ("mh_f", "Mean hourly pay, women", "avg", "hourly_total", "F", False),
            ("mh_m", "Mean hourly pay, men", "avg", "hourly_total", "M", False),
            ("ma_f", "Mean annual pay, women", "avg", "total_pay", "F", False),
            ("ma_m", "Mean annual pay, men", "avg", "total_pay", "M", False),
            ("dh_f", "Median hourly pay, women", "med", "hourly_total", "F", False),
            ("dh_m", "Median hourly pay, men", "med", "hourly_total", "M", False),
            ("da_f", "Median annual pay, women", "med", "total_pay", "F", False),
            ("da_m", "Median annual pay, men", "med", "total_pay", "M", False),
            ("rc_f", "Women receiving complementary or variable components", "recv", "F"),
            ("rc_m", "Men receiving complementary or variable components", "recv", "M"),
            ("mc_f", "Mean components among recipients, women", "avg", "components", "F", True),
            ("mc_m", "Mean components among recipients, men", "avg", "components", "M", True),
            ("dc_f", "Median components among recipients, women", "med", "components", "F", True),
            ("dc_m", "Median components among recipients, men", "med", "components", "M", True),
            ("mb_f", "Mean hourly base pay, women", "avg", "hourly_base", "F", False),
            ("mb_m", "Mean hourly base pay, men", "avg", "hourly_base", "M", False),
        ]
        for q in (1, 2, 3, 4):
            specs.append((f"q{q}_n", f"Workers in quartile pay band {q}", "q", q, None))
            specs.append((f"q{q}_f", f"Women in quartile pay band {q}", "q", q, "F"))
        fmts = {"count": COUNT, "recv": COUNT, "q": COUNT}
        for spec in specs:
            key, label, kind = spec[:3]
            work[key] = r
            put(ws, f"B{r}", label)
            for code, L in cols.items():
                group = code == "Group"
                ctry = "" if group else f"{c('country')},{L}$4,"
                cond = "" if group else f"({c('country')}={L}$4)*"
                ref = f"{L}{r}"
                if kind == "count":
                    if key == "workers":
                        f = f"=COUNTA({c('worker_id')})" if group else f"=COUNTIF({c('country')},{L}$4)"
                    else:
                        sex = "F" if key == "women" else "M"
                        f = f'=COUNTIFS({ctry}{c("gender")},"{sex}")'
                    put(ws, ref, f, fmt=COUNT)
                elif kind == "avg":
                    _, _, _, field, sex, recipients = spec
                    extra = f",{c('receives_components')},1" if recipients else ""
                    fmt = EUR_H if field.startswith("hourly") else EUR
                    put(ws, ref, f'=IFERROR(AVERAGEIFS({c(field)},{ctry}{c("gender")},"{sex}"{extra}),"")', fmt=fmt)
                elif kind == "med":
                    _, _, _, field, sex, recipients = spec
                    extra = f"*({c('receives_components')}=1)" if recipients else ""
                    fmt = EUR_H if field.startswith("hourly") else EUR
                    array(ws, ref, f'=IFERROR(MEDIAN(IF({cond}({c("gender")}="{sex}"){extra},{c(field)})),"")', fmt)
                elif kind == "recv":
                    sex = spec[3]
                    put(ws, ref, f'=COUNTIFS({ctry}{c("gender")},"{sex}",{c("receives_components")},1)', fmt=COUNT)
                elif kind == "q":
                    q, sex = spec[3], spec[4]
                    sexpart = f',{c("gender")},"{sex}"' if sex else ""
                    put(ws, ref, f"=COUNTIFS({ctry}{c('quartile')},{q}{sexpart})", fmt=COUNT)
            r += 1
        self.work_rows = work

        # Reporting items
        W = work
        definitions = {
            "workers": "Head count at the snapshot, used for the size band.",
            "a_mean_gap_hourly": "Mean gross hourly pay, base plus all components.",
            "a_mean_gap_annual": "Mean gross annual pay, base plus all components.",
            "b_mean_gap_components": "Mean annual bonus, commission and allowances, among workers who receive any.",
            "c_median_gap_hourly": "Median gross hourly pay.",
            "c_median_gap_annual": "Median gross annual pay.",
            "d_median_gap_components": "Median annual components, among workers who receive any.",
            "e_share_women_receiving": "Share of women with any complementary or variable component.",
            "f_q1_share_women": "Four equal groups by hourly pay within the employer (Article 3(1)(f)).",
            "base_mean_gap_hourly": "Not a reporting item: reconciles with the 15.4% base-pay gap in hr-people-analytics.",
        }
        formulas = {
            "workers": lambda L: f"={L}{W['workers']}",
            "women": lambda L: f"={L}{W['women']}",
            "men": lambda L: f"={L}{W['men']}",
            "a_mean_gap_hourly": lambda L: f'=IFERROR(({L}{W["mh_m"]}-{L}{W["mh_f"]})/{L}{W["mh_m"]},"")',
            "a_mean_gap_annual": lambda L: f'=IFERROR(({L}{W["ma_m"]}-{L}{W["ma_f"]})/{L}{W["ma_m"]},"")',
            "b_mean_gap_components": lambda L: f'=IFERROR(({L}{W["mc_m"]}-{L}{W["mc_f"]})/{L}{W["mc_m"]},"")',
            "c_median_gap_hourly": lambda L: f'=IFERROR(({L}{W["dh_m"]}-{L}{W["dh_f"]})/{L}{W["dh_m"]},"")',
            "c_median_gap_annual": lambda L: f'=IFERROR(({L}{W["da_m"]}-{L}{W["da_f"]})/{L}{W["da_m"]},"")',
            "d_median_gap_components": lambda L: f'=IFERROR(({L}{W["dc_m"]}-{L}{W["dc_f"]})/{L}{W["dc_m"]},"")',
            "e_share_women_receiving": lambda L: f'=IFERROR({L}{W["rc_f"]}/{L}{W["women"]},"")',
            "e_share_men_receiving": lambda L: f'=IFERROR({L}{W["rc_m"]}/{L}{W["men"]},"")',
            "base_mean_gap_hourly": lambda L: f'=IFERROR(({L}{W["mb_m"]}-{L}{W["mb_f"]})/{L}{W["mb_m"]},"")',
        }
        for q in (1, 2, 3, 4):
            formulas[f"f_q{q}_share_women"] = (lambda q: lambda L: f'=IFERROR({L}{W[f"q{q}_f"]}/{L}{W[f"q{q}_n"]},"")')(q)
            formulas[f"f_q{q}_share_men"] = (lambda q: lambda L: f'=IFERROR(({L}{W[f"q{q}_n"]}-{L}{W[f"q{q}_f"]})/{L}{W[f"q{q}_n"]},"")')(q)

        python = article9.by_entity(self.w, codes)
        r = 6
        self.item_rows = {}
        for key, label, kind in article9.ITEMS:
            self.item_rows[key] = r
            is_context = key.startswith("base_")
            put(ws, f"B{r}", label, italic=is_context, color=MUTED if is_context else INK)
            for code, L in cols.items():
                put(ws, f"{L}{r}", formulas[key](L), fmt=COUNT if kind == "count" else PCT,
                    bold=key.startswith(("a_", "c_")))
                self.check("Article 9", code, label, python.loc[key, code], f"'Article 9'!{L}{r}")
            put(ws, f"H{r}", definitions.get(key, ""), color=MUTED, italic=True)
            if r % 2 == 0:
                for L in ["B", *cols.values()]:
                    ws[f"{L}{r}"].fill = BAND_FILL
            r += 1
        for code, L in cols.items():
            for q in (1, 2, 3, 4):
                self.check("Article 9 workings", code, f"Workers in quartile pay band {q}",
                           python.loc[f"f_q{q}_workers", code], f"'Article 9'!{L}{W[f'q{q}_n']}")
        ws.freeze_panes = "C6"
        self.cols9 = cols

    def categories(self, ws):
        c = self.col
        title(ws, "Article 9(1)(g) by category of workers, and the Article 10 screen",
              "A category at or above the threshold on hourly or annual pay needs an objective, gender-neutral "
              "justification, or a remedy within six months of the report. Otherwise a joint pay assessment follows.")
        heads = ["Employer", "Category", "Women", "Men",
                 "Hourly base, women", "Hourly base, men", "Gap: base",
                 "Hourly components, women", "Hourly components, men", "Gap: components",
                 "Hourly pay, women", "Hourly pay, men", "Gap: hourly pay",
                 "Annual pay, women", "Annual pay, men", "Gap: annual pay",
                 "Fewer than the minimum of one sex", "Screen result",
                 "Justification (objective, gender-neutral)", "Remedied within six months?",
                 "Article 10 outcome", "Cost to close the hourly gap, upper bound (€ a year)"]
        header(ws, 4, heads, height=58)
        widths(ws, {get_column_letter(i): w for i, w in enumerate(
            [10, 10, 8, 8, 10, 10, 9, 11, 11, 11, 10, 10, 10, 11, 11, 10, 12, 20, 34, 13, 30, 16], start=1)})
        python = article10.screen(self.w, self.countries, [b for _, b in BANDS]).set_index(["country", "category"])
        thr, mincell = self.refs["threshold"], self.refs["min_cell"]
        r = 5
        crit = lambda r_, sex: f'{c("country")},$A{r_},{c("category")},$B{r_},{c("gender")},"{sex}"'
        dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
        ws.add_data_validation(dv)
        first = r
        for code in self.countries:
            for _, cat in BANDS:
                put(ws, f"A{r}", code, color=BLUE)
                put(ws, f"B{r}", cat, color=BLUE)
                put(ws, f"C{r}", f"=COUNTIFS({crit(r, 'F')})", fmt=COUNT)
                put(ws, f"D{r}", f"=COUNTIFS({crit(r, 'M')})", fmt=COUNT)
                pairs = [("E", "F", "G", "hourly_base", EUR_H), ("H", "I", "J", "hourly_components", EUR_H),
                         ("K", "L", "M", "hourly_total", EUR_H), ("N", "O", "P", "total_pay", EUR)]
                for fcol, mcol, gcol, field, fmt in pairs:
                    put(ws, f"{fcol}{r}", f'=IFERROR(AVERAGEIFS({c(field)},{crit(r, "F")}),"")', fmt=fmt)
                    put(ws, f"{mcol}{r}", f'=IFERROR(AVERAGEIFS({c(field)},{crit(r, "M")}),"")', fmt=fmt)
                    put(ws, f"{gcol}{r}", f'=IFERROR(({mcol}{r}-{fcol}{r})/{mcol}{r},"")', fmt=PCT, bold=field == "hourly_total")
                put(ws, f"Q{r}", f'=IF(MIN(C{r},D{r})<{mincell},"Yes","No")')
                put(ws, f"R{r}", (f'=IF(C{r}+D{r}=0,"No workers",IF(OR(C{r}=0,D{r}=0),"{article10.STATUS_ONE_SEX}",'
                                  f'IF(OR(ABS(M{r})>={thr},ABS(P{r})>={thr}),"{article10.STATUS_FLAG}","{article10.STATUS_OK}")))'))
                put(ws, f"S{r}", None, color=BLUE)
                put(ws, f"T{r}", None, color=BLUE)
                dv.add(f"T{r}")
                put(ws, f"U{r}", (f'=IF(R{r}<>"{article10.STATUS_FLAG}","Not required",IF(LEN(S{r})>0,'
                                  f'"Record the justification and share it",IF(T{r}="Yes","Remedied",'
                                  f'"Joint pay assessment unless remedied in time")))'), wrap=True)
                put(ws, f"V{r}", (f'=IF(R{r}="{article10.STATUS_FLAG}",IF(M{r}>0,(L{r}-K{r})*SUMIFS({c("annual_hours")},{crit(r, "F")}),'
                                  f'(K{r}-L{r})*SUMIFS({c("annual_hours")},{crit(r, "M")})),0)'), fmt=EUR)
                if (code, cat) in python.index:
                    p = python.loc[(code, cat)]
                    label = f"{cat}"
                    for item, col_ in [("women", "C"), ("men", "D"), ("gap_base", "G"), ("gap_components", "J"),
                                       ("gap_hourly", "M"), ("gap_annual", "P"), ("status", "R"),
                                       ("remedy_cost", "V")]:
                        value = p[item]
                        self.check("Categories", code, f"{label}: {item.replace('_', ' ')}",
                                   value.item() if hasattr(value, "item") else value, f"Categories!{col_}{r}")
                    self.check("Categories", code, f"{label}: small cell",
                               "Yes" if p["small_cell"] else "No", f"Categories!Q{r}")
                r += 1
        last = r - 1
        self.cat_range = (first, last)
        status_colours(ws, f"R{first}:R{last}")
        ws.conditional_formatting.add(f"Q{first}:Q{last}", CellIsRule(
            operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor=STATUS_FILLS["Partial"])))
        r += 1
        put(ws, f"B{r}", "Total", bold=True)
        put(ws, f"R{r}", f'=COUNTIF(R{first}:R{last},"{article10.STATUS_FLAG}")&" to justify or remedy"', bold=True)
        put(ws, f"V{r}", f"=SUM(V{first}:V{last})", bold=True, fmt=EUR)
        self.check("Categories", "", "Total cost to close, upper bound",
                   float(python["remedy_cost"].sum()), f"Categories!V{r}")
        r += 2
        notes = [
            "Components per hour = bonus, commission and allowances divided by annual contractual hours, averaged over all workers in the category.",
            "The cost column lifts the lower-paid sex's average hourly pay to the other's, for all their hours. It is a ceiling: any justified part of a gap reduces it.",
            "Where one sex has fewer workers than the minimum in Settings, share the figures only with workers' representatives, the labour inspectorate or the equality body (Article 12(3)).",
        ]
        for note in notes:
            put(ws, f"A{r}", note, color=MUTED, italic=True)
            r += 1
        ws.freeze_panes = "C5"
        ws.auto_filter.ref = f"A4:V{last}"

    def obligations(self, ws):
        c = self.col
        title(ws, "What each employer owes, and when",
              f"Legal status as checked on {ASSESSMENT_DATE:%d %B %Y}. Items marked [to verify] need the final national text.")
        heads = ["Code", "Employer", "Workers", "Size band", "Reporting frequency", "First report due",
                 "Remedy deadline if a category is flagged", "Transposition", "National instrument",
                 "National reporting rules", "Notes", "Sources", "Checked on"]
        header(ws, 4, heads, height=44)
        widths(ws, {"A": 7, "B": 16, "C": 9, "D": 14, "E": 18, "F": 13, "G": 16, "H": 16, "I": 42,
                    "J": 42, "K": 60, "L": 44, "M": 12})
        status = obligations.transposition().set_index("country")
        python = obligations.by_entity(self.w).set_index("country")
        self.obl_rows = {}
        for r, code in enumerate(self.countries, start=5):
            self.obl_rows[code] = r
            put(ws, f"A{r}", code, color=BLUE)
            put(ws, f"B{r}", ENTITIES[code], color=BLUE)
            put(ws, f"C{r}", f"=COUNTIF({c('country')},A{r})", fmt=COUNT)
            match = f"MATCH(C{r},{self.refs['size_floor']},1)"
            put(ws, f"D{r}", f"=INDEX({self.refs['size_label']},{match})")
            put(ws, f"E{r}", f"=INDEX({self.refs['size_freq']},{match})")
            put(ws, f"F{r}", f"=INDEX({self.refs['size_first']},{match})", fmt=DATE, bold=True)
            put(ws, f"G{r}", f'=IFERROR(EDATE(F{r},{self.refs["months"]}),"n/a")', fmt=DATE)
            s = status.loc[code]
            put(ws, f"H{r}", s["status"], color=BLUE, bold=True)
            for col_, field in zip("IJKL", ["instrument", "reporting", "notes", "sources"]):
                put(ws, f"{col_}{r}", s[field], color=BLUE, wrap=True)
            put(ws, f"M{r}", s["checked_on"], color=BLUE)
            ws.row_dimensions[r].height = 92
            p = python.loc[code]
            self.check("Obligations", code, "workers", int(p["workers"]), f"Obligations!C{r}")
            self.check("Obligations", code, "size band", p["size_band"], f"Obligations!D{r}")
        for text, colour in {"Transposed": "D5ECDC", "Partly transposed": "FBE8B5",
                             "Draft": "FBE8B5", "Not transposed": "F5C6C2"}.items():
            ws.conditional_formatting.add(f"H5:H{4 + len(self.countries)}", CellIsRule(operator="equal", formula=[f'"{text}"'],
                                                              fill=PatternFill("solid", fgColor=colour)))
        put(ws, "A11", "Not legal advice. The national act, and any ministry guidance on method, decide how the "
                       "report is finally computed and filed.", color=MUTED, italic=True)

    def readiness(self, ws):
        title(ws, "Readiness: what has to be in place, and how far along it is",
              "Statuses are illustrative for the demonstration organisation. Change them with the drop-down.")
        widths(ws, {"A": 6, "B": 16, "C": 62, "D": 44, "E": 11, "F": 18, "G": 52})
        items = obligations.readiness()
        first, last = 8, 7 + len(items)
        put(ws, "B4", "Ready", bold=True)
        put(ws, "B5", "Partial", bold=True)
        put(ws, "B6", "Gap", bold=True)
        for i, s in enumerate(["Ready", "Partial", "Gap"]):
            put(ws, f"C{4 + i}", f'=COUNTIF($E${first}:$E${last},"{s}")', fmt="0", align="left")
            self.refs[f"ready_{s}"] = f"Readiness!$C${4 + i}"
        status_colours(ws, "B4:B6")
        header(ws, 7, ["Ref", "Article", "Requirement", "Evidence expected", "Status", "Owner", "Note"])
        dv = DataValidation(type="list", formula1='"Ready,Partial,Gap"', allow_blank=False)
        ws.add_data_validation(dv)
        for r, it in enumerate(items.itertuples(index=False), start=first):
            for col_, value in zip("ABCDEFG", [it.ref, it.article, it.requirement, it.evidence, it.status,
                                               it.owner, it.note]):
                put(ws, f"{col_}{r}", value, color=BLUE, wrap=True)
            dv.add(f"E{r}")
            ws.row_dimensions[r].height = 44
        status_colours(ws, f"E{first}:E{last}")
        counts = items["status"].value_counts()
        for i, s in enumerate(["Ready", "Partial", "Gap"]):
            self.check("Readiness", "", f"{s} count", int(counts.get(s, 0)), f"Readiness!C{4 + i}")
        ws.freeze_panes = "C8"

    def summary(self, ws):
        title(ws, "Pay transparency readiness: summary",
              f"Dry run on synthetic data for workers active on {SNAPSHOT:%d %B %Y}. Not a statutory submission.")
        widths(ws, {"A": 3, "B": 50, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14})
        codes = self.countries
        header(ws, 4, ["", "Measure", *codes, "Group"])
        cols = self.cols9
        put(ws, "B5", "Employer", italic=True, color=MUTED)
        for code in codes:
            put(ws, f"{cols[code]}5", f"='Article 9'!{cols[code]}5", color=GREEN, italic=True)
        first_cat, last_cat = self.cat_range
        rows = [
            ("Workers", "workers", COUNT),
            ("Gender pay gap, mean hourly pay", "a_mean_gap_hourly", PCT),
            ("Median gender pay gap, hourly pay", "c_median_gap_hourly", PCT),
            ("Gap in complementary or variable components, mean", "b_mean_gap_components", PCT),
            ("Women in the highest quartile pay band", "f_q4_share_women", PCT),
        ]
        r = 6
        for label, key, fmt in rows:
            put(ws, f"B{r}", label, bold=key.startswith("a_"))
            for code, L in cols.items():
                put(ws, f"{L}{r}", f"='Article 9'!{L}{self.item_rows[key]}", color=GREEN, fmt=fmt,
                    bold=key.startswith("a_"))
            r += 1
        self.summary_gap_rows = (7, 8)
        put(ws, f"B{r}", "First report due")
        for code in codes:
            put(ws, f"{cols[code]}{r}", f"=Obligations!F{self.obl_rows[code]}", color=GREEN, fmt=DATE)
        put(ws, f"{cols['Group']}{r}", f"=MIN({cols[codes[0]]}{r}:{cols[codes[-1]]}{r})", fmt=DATE)
        r += 1
        put(ws, f"B{r}", "Transposition status")
        for code in codes:
            put(ws, f"{cols[code]}{r}", f"=Obligations!H{self.obl_rows[code]}", color=GREEN)
        r += 1
        cat_rows = r
        put(ws, f"B{r}", "Categories to justify or remedy (5% or more)", bold=True)
        put(ws, f"B{r + 1}", "Categories with too few workers of one sex to share widely")
        put(ws, f"B{r + 2}", "Cost to close flagged hourly gaps, upper bound (€ a year)", bold=True)
        for code in codes:
            L = cols[code]
            put(ws, f"{L}{r}", (f'=COUNTIFS(Categories!$A${first_cat}:$A${last_cat},{L}$4,'
                                f'Categories!$R${first_cat}:$R${last_cat},"{article10.STATUS_FLAG}")'), bold=True)
            put(ws, f"{L}{r + 1}", (f'=COUNTIFS(Categories!$A${first_cat}:$A${last_cat},{L}$4,'
                                    f'Categories!$Q${first_cat}:$Q${last_cat},"Yes")'))
            put(ws, f"{L}{r + 2}", (f"=SUMIFS(Categories!$V${first_cat}:$V${last_cat},"
                                    f"Categories!$A${first_cat}:$A${last_cat},{L}$4)"), fmt=EUR, bold=True)
        G, a, b = cols["Group"], cols[codes[0]], cols[codes[-1]]
        for k in range(3):
            put(ws, f"{G}{r + k}", f"=SUM({a}{r + k}:{b}{r + k})", fmt=EUR if k == 2 else "0", bold=k != 1)
        python = article10.screen(self.w, codes, [b_ for _, b_ in BANDS])
        for code in codes:
            sub = python[python["country"] == code]
            self.check("Summary", code, "categories flagged", int((sub["status"] == article10.STATUS_FLAG).sum()),
                       f"Summary!{cols[code]}{r}")
        r += 4
        put(ws, f"B{r}", "Readiness checklist (Articles 4 to 12)", bold=True, size=11)
        for i, s in enumerate(["Ready", "Partial", "Gap"]):
            put(ws, f"B{r + 1 + i}", s)
            put(ws, f"C{r + 1 + i}", f"={self.refs[f'ready_{s}']}", color=GREEN, fmt="0")
        status_colours(ws, f"B{r + 1}:B{r + 3}")

        chart = BarChart()
        chart.type = "bar"
        chart.title = "Gender pay gap by employer, hourly pay"
        chart.y_axis.numFmt = "0%"
        chart.y_axis.majorGridlines = None
        data = Reference(ws, min_col=2, max_col=2 + len(codes), min_row=7, max_row=8)
        chart.add_data(data, from_rows=True, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=3, max_col=2 + len(codes), min_row=4, max_row=4))
        chart.height, chart.width = 7.5, 16
        ws.add_chart(chart, "I4")

        chart2 = BarChart()
        chart2.type = "col"
        chart2.title = "Categories to justify or remedy"
        chart2.legend = None
        chart2.y_axis.majorGridlines = None
        data = Reference(ws, min_col=2, max_col=2 + len(codes), min_row=cat_rows, max_row=cat_rows)
        chart2.add_data(data, from_rows=True, titles_from_data=True)
        chart2.set_categories(Reference(ws, min_col=3, max_col=2 + len(codes), min_row=4, max_row=4))
        chart2.height, chart2.width = 7.5, 16
        ws.add_chart(chart2, "I20")

    def reconciliation(self, ws):
        title(ws, "Reconciliation: workbook formulas against the Python pipeline",
              "Column D was written by src/build_workbook.py from the pandas results; column E is the live formula. "
              "Numbers must agree to one part in a billion. After pasting a new extract, rerun the pipeline to refresh D.")
        widths(ws, {"A": 18, "B": 8, "C": 60, "D": 18, "E": 18, "F": 14, "G": 8})
        header(ws, 6, ["Area", "Employer", "Item", "Python value", "Workbook value", "Difference", "Match"])
        first = 7
        last = first + len(self.checks) - 1
        put(ws, "B3", "Result", bold=True)
        put(ws, "C3", (f'=IF(COUNTIF(G{first}:G{last},"No")=0,"All "&COUNTA(C{first}:C{last})&" checks match",'
                       f'COUNTIF(G{first}:G{last},"No")&" of "&COUNTA(C{first}:C{last})&" checks do not match")'),
            bold=True, size=12)
        for r, (area, entity, item, value, ref) in enumerate(self.checks, start=first):
            put(ws, f"A{r}", area)
            put(ws, f"B{r}", entity)
            put(ws, f"C{r}", item)
            fmt = "0.000000000" if isinstance(value, float) else None
            put(ws, f"D{r}", value, color=BLUE, fmt=fmt)
            put(ws, f"E{r}", f"={ref}", color=GREEN, fmt=fmt)
            put(ws, f"F{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"")', fmt="0.0E+00")
            put(ws, f"G{r}", (f'=IF(AND(D{r}="",E{r}=""),"Yes",IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),'
                              f'IF(ABS(E{r}-D{r})<=1E-9*MAX(1,ABS(D{r})),"Yes","No"),IF(D{r}=E{r},"Yes","No")))'))
        status_colours(ws, f"G{first}:G{last}")
        ws.freeze_panes = "A7"
        self.recon_result = "Reconciliation!C3"

    def cover(self, ws, names):
        widths(ws, {"A": 3, "B": 30, "C": 100})
        put(ws, "B2", "Pay Transparency Readiness Model", bold=True, size=18)
        put(ws, "B3", "Directive (EU) 2023/970 - a dry run for an employer in four Member States", color=MUTED, size=12)
        put(ws, "B5", "What it does", bold=True, size=11)
        lines = [
            "Computes the seven Article 9(1) reporting items for each employer from a payroll extract.",
            "Scores every role on the four Article 4(4) criteria and groups workers into categories of equal value.",
            "Screens each category against the 5% threshold of Article 10 and prices the cost of closing the gaps.",
            "Sets out each employer's deadlines, the state of national transposition, and a readiness checklist.",
        ]
        for i, line in enumerate(lines, start=6):
            put(ws, f"C{i}", line)
        put(ws, "B11", "How to read it", bold=True, size=11)
        legend = [
            ("Blue text", "an input: data, a score, an assumption", BLUE, None),
            ("Black text", "a formula", INK, None),
            ("Green text", "a link to another sheet", GREEN, None),
            ("Yellow fill", "a key assumption that moves the results", INK, YELLOW),
        ]
        for i, (label, meaning, colour, fill) in enumerate(legend, start=12):
            put(ws, f"B{i}", label, color=colour, fill=fill)
            put(ws, f"C{i}", meaning)
        put(ws, "B17", "Sheets", bold=True, size=11)
        purpose = {
            "Summary": "The page for the board: headline gaps, flagged categories, cost, readiness.",
            "Article 9": "The seven reporting items for each employer, with the workings underneath.",
            "Categories": "Pay gaps by category of workers and the Article 10 screen, with space for justifications.",
            "Obligations": "Size band, deadlines and national transposition status for each employer.",
            "Readiness": "Fourteen obligations from Articles 4 to 12, with evidence, owner and status.",
            "Job Evaluation": "The scored role grid that defines the categories of workers.",
            "Payroll": "One row per worker: the extract (blue) and the derived pay figures (black).",
            "Settings": "Every assumption: hours, thresholds, factor weights, point bands, size bands.",
            "Reconciliation": "Each workbook figure checked against the Python pipeline.",
        }
        for i, n in enumerate(names[1:], start=18):
            cell = put(ws, f"B{i}", n, color="1F5FA8")
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{n}'!A1")
            put(ws, f"C{i}", purpose[n])
        put(ws, "B28", "Check", bold=True, size=11)
        put(ws, "C28", f"={self.recon_result}", color=GREEN, bold=True)
        put(ws, "B30", "Data", bold=True, size=11)
        put(ws, "C30", "Synthetic. The organisation is the one analysed in hr-people-analytics; the pay components were "
                       "added by this kit with rules stated in data/README.md. No real person's pay is in this file.", wrap=True)
        ws.row_dimensions[30].height = 28
        put(ws, "B31", "Your own data", bold=True, size=11)
        put(ws, "C31", "Paste an extract with the same columns into Payroll A:R, extend the black columns to the last row, "
                       "and rerun the Python pipeline so the Reconciliation sheet compares against your numbers.", wrap=True)
        ws.row_dimensions[31].height = 28
        put(ws, "B32", "Limits", bold=True, size=11)
        put(ws, "C32", "Not legal advice. National transposition decides the final method and filing; see Obligations "
                       "and docs/verification.md in the repository.", wrap=True)
        put(ws, "B34", "Domenico Perroni", color=MUTED)
        cell = put(ws, "C34", REPO, color="1F5FA8")
        cell.hyperlink = REPO


def build(workers: pd.DataFrame | None = None, path: Path = OUTPUT) -> Model:
    workers = load_workers() if workers is None else workers
    model = Model(workers, load_job_evaluation())
    model.build(path)
    return model


def main() -> None:
    model = build()
    print(f"workbook: {len(model.checks)} reconciliation checks -> {OUTPUT.relative_to(OUTPUT.parents[1])}")


if __name__ == "__main__":
    main()
